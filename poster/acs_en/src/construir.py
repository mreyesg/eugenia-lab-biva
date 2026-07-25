# Construye el poster de ACS en ingles sobre la plantilla de Manuel (36 x 48 pulgadas),
# conservando su paleta y su estructura de columnas. El contenido factual sale del poster
# en espanol y de CLAUDE.md; las figuras son capturas reales de produccion.
import copy, io, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

PLANTILLA = "plantilla.pptx"
SALIDA = "ACS_Poster_EugenIA_Lab_EN.pptx"

# Paleta leida de la plantilla
TINTA   = RGBColor(0x1B, 0x1B, 0x1B)
MORADO  = RGBColor(0x6A, 0x1B, 0x9A)
TEAL    = RGBColor(0x4E, 0xC1, 0xB9)
GRIS    = RGBColor(0x55, 0x55, 0x55)
PANEL   = RGBColor(0xFF, 0xFF, 0xFF)
BORDE   = RGBColor(0xEA, 0xE7, 0xDF)
SUAVE   = RGBColor(0xF4, 0xF1, 0xEA)
CODIGO  = RGBColor(0x1E, 0x29, 0x3B)

prs = Presentation(PLANTILLA)
sl  = prs.slides[0]

# ---------- utilidades ----------
def por_id(i):
    for sh in sl.shapes:
        if sh.shape_id == i: return sh
    return None

def borrar(*ids):
    for i in ids:
        sh = por_id(i)
        if sh is not None: sh._element.getparent().remove(sh._element)

def escribir(sh, bloques, tam=32, color=TINTA, bold=False, interlineado=0.92, espacio=10,
             align=None, fuente=None):
    """bloques: str o lista de (texto, dict de overrides)"""
    tf = sh.text_frame
    tf.word_wrap = True
    tf.clear()
    if isinstance(bloques, str): bloques = [(bloques, {})]
    for k, item in enumerate(bloques):
        txt, ov = item if isinstance(item, tuple) else (item, {})
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.line_spacing = ov.get("interlineado", interlineado)
        p.space_after = Pt(ov.get("espacio", espacio))
        al = ov.get("align", align)
        if al is not None: p.alignment = al
        r = p.add_run(); r.text = txt
        f = r.font
        f.size  = Pt(ov.get("tam", tam))
        f.bold  = ov.get("bold", bold)
        f.color.rgb = ov.get("color", color)
        f.name  = ov.get("fuente", fuente or "Calibri")
    return sh

def caja(x, y, w, h, texto, **kw):
    sh = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    escribir(sh, texto, **kw)
    return sh

def encabezado(x, y, w, texto, tam=54, color=MORADO):
    return caja(x, y, w, 0.95, texto, tam=tam, color=color, bold=True, espacio=0)

def panel(x, y, w, h, relleno=PANEL, borde=BORDE):
    from pptx.enum.shapes import MSO_SHAPE
    sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = relleno
    sh.line.color.rgb = borde; sh.line.width = Pt(1.5)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh

# Las divisorias intermedias (y=9,15,21,27,33,39) cruzaban el texto como tachones.
borrar(42, 43, 44, 45, 46, 47)

# ---------- HEADER ----------
escribir(por_id(8),
    "AI-Assisted Literature Analysis for the Experimental Design of Solvent Optimization "
    "in Analgesic Studies of Plectranthus amboinicus",
    tam=54, bold=True, interlineado=0.9, espacio=0)
escribir(por_id(9),
    [("Kamila López Díaz  ·  Manuel Reyes-Guzmán, Ph.D. (principal investigator)", {"tam": 34, "bold": True}),
     ("EugenIA Lab  ·  Department of Physical Sciences (CIFI 4996)  ·  Faculty of General Studies  ·  "
      "University of Puerto Rico, Río Piedras", {"tam": 27, "color": GRIS})],
    interlineado=0.95, espacio=3)
por_id(8).width = Inches(28.0); por_id(8).height = Inches(1.7)
por_id(9).width = Inches(28.5); por_id(9).top = Inches(1.95); por_id(9).height = Inches(1.0)

# Logo institucional. Es cuadrado (900x900), asi que se coloca a escala 1:1
# dentro del recuadro de logos de la plantilla, sin deformarlo.
LOGO = "logo_upr.jpg"
if os.path.exists(LOGO):
    borrar(11)  # el marcador de texto "UPR-RP LOGO"
    lado = 1.9
    sl.shapes.add_picture(LOGO, Inches(30.75), Inches(0.7), Inches(lado), Inches(lado))
    ac = por_id(12)  # "ACS PR LOGO", se corre al lado para no quedar debajo del logo
    if ac is not None:
        ac.left, ac.top, ac.width = Inches(32.85), Inches(1.25), Inches(2.4)
        escribir(ac, "ACS PR LOGO", tam=20, color=GRIS, espacio=0)

# ---------- BARRA IZQUIERDA ----------
escribir(por_id(13), "Overview", tam=46, color=MORADO, bold=True, espacio=0)
escribir(por_id(14),
    [("Aim: turn curated literature into concrete experimental design criteria.", {}),
     ("Species: Plectranthus amboinicus (Cuban oregano).", {}),
     ("Agent: EugenIA, a multi-database bibliometric pipeline.", {}),
     ("Artifact: the Experimental Decision Matrix (EDM).", {}),
     ("Question: which ethanol concentration maximizes total phenolic content?", {})],
    tam=23, interlineado=0.93, espacio=7)
por_id(14).left = Inches(0.85); por_id(14).width = Inches(4.3); por_id(14).height = Inches(5.0)

escribir(por_id(15), "Navigation", tam=46, color=MORADO, bold=True, espacio=0)
por_id(15).top = Inches(10.1); por_id(15).left = Inches(0.85); por_id(15).width = Inches(4.3)
escribir(por_id(16),
    ["1.  Abstract", "2.  Methods: the EugenIA pipeline", "3.  From literature to decision",
     "4.  The Decision Matrix", "5.  Results and figures", "6.  Discussion", "7.  Conclusions"],
    tam=23, interlineado=0.95, espacio=5)
por_id(16).left = Inches(0.85); por_id(16).width = Inches(4.3); por_id(16).top = Inches(11.0); por_id(16).height = Inches(3.6)

escribir(por_id(17), "Try the tool", tam=46, color=MORADO, bold=True, espacio=0)
por_id(17).top = Inches(14.6); por_id(17).left = Inches(0.85); por_id(17).width = Inches(4.3)

# QR real de produccion
qr_marco = por_id(18)
qx, qy = Inches(0.9), Inches(15.6)
qs = Inches(3.5)
qr_marco.left, qr_marco.top, qr_marco.width, qr_marco.height = qx, qy, qs, qs
QR = "H:/VS Code Projects/eugenia_lab_biva/poster/assets/qr-produccion.png"
if os.path.exists(QR):
    sl.shapes.add_picture(QR, qx + Emu(60000), qy + Emu(60000), qs - Emu(120000), qs - Emu(120000))
borrar(19)
caja(0.85, 19.3, 4.3, 2.1,
     [("Biblioteca Viva", {"tam": 28, "bold": True, "color": MORADO}),
      ("Scan to run this workflow on your own phone. No install, no account.", {"tam": 24, "color": GRIS})],
     interlineado=0.95, espacio=6)
escribir(por_id(20),
    [("mreyes@upr.edu", {"tam": 24, "bold": True}),
     ("web-beta-hazel-49.vercel.app", {"tam": 22, "color": GRIS})],
    interlineado=0.95, espacio=3)
por_id(20).top = Inches(21.5); por_id(20).left = Inches(0.85); por_id(20).width = Inches(4.3)

# ---- At a glance: numeros grandes en la barra ----
encabezado(0.85, 23.3, 4.3, "At a glance", tam=46)
cifras = [
    ("47",      "curated articles in the corpus"),
    ("20–40%",  "v/v aqueous ethanol, most reported solvent"),
    ("20–80",   "mg GAE/g, total phenolic content in polar solvents"),
    ("91.3%",   "agreement, automated vs manual extraction"),
    ("50–70%",  "v/v ethanol, criterion adopted for the teaching lab"),
    ("3",       "databases searched: OpenAlex, Europe PMC, Crossref"),
]
yy = 24.4
for num, txt in cifras:
    panel(0.8, yy, 4.5, 3.0, relleno=SUAVE, borde=BORDE)
    caja(1.0, yy + 0.15, 4.1, 1.3, num, tam=60, color=MORADO, bold=True, espacio=0, interlineado=0.9)
    caja(1.0, yy + 1.45, 4.1, 1.4, txt, tam=22, color=GRIS, espacio=0, interlineado=0.92)
    yy += 3.25

# ---------- COLUMNA CENTRAL ----------
escribir(por_id(21), "Abstract", tam=54, color=MORADO, bold=True, espacio=0)
escribir(por_id(22),
    "Undergraduate research in phytochemistry often stalls between reading the literature and "
    "choosing experimental conditions. This project asks whether an AI bibliometric agent can close "
    "that gap for Plectranthus amboinicus, a species with abundant but methodologically heterogeneous "
    "literature. EugenIA, the agent built at EugenIA Lab, searches multiple bibliographic databases, "
    "deduplicates by DOI and normalized title, validates open-access PDFs and produces a curated corpus. "
    "That corpus feeds an Experimental Decision Matrix (EDM): a typed artifact pairing what the "
    "literature reports with the operating criterion the teaching laboratory will adopt, and recording "
    "which curated papers support each row. Across a 47-article corpus, aqueous ethanol (20–40% v/v) and "
    "methanol were the most reported solvents, with total phenolic content between 20 and 80 mg GAE/g; "
    "Soxhlet and maceration dominated. Agreement between automated and manual variable extraction was 91.3%.",
    tam=28, interlineado=0.93, espacio=0)
por_id(22).top = Inches(4.6); por_id(22).height = Inches(6.0)

escribir(por_id(23), "Methods: the EugenIA pipeline", tam=54, color=MORADO, bold=True, espacio=0)
por_id(23).top = Inches(10.8)
escribir(por_id(24),
    [("EugenIA runs a reproducible pipeline, not a chat.", {"bold": True}),
     ("Search across OpenAlex, Europe PMC and Crossref, with the executed query shown to the researcher and editable.", {}),
     ("Deduplication by DOI and normalized title, then interleaving of the three sources.", {}),
     ("Open-access enrichment and PDF validation.", {}),
     ("Structured annotation of every retrieved article: its contribution, and how strongly it supports the component being built.", {}),
     ("Export of the curated corpus to BibTeX, RIS and CSV.", {}),
     ("Every article shown comes from a database, never from the model. Each search is logged with its query, "
      "filters and per-database counts, and the report renders a methods paragraph ready to paste into a manuscript.",
      {"color": GRIS, "tam": 28})],
    tam=28, interlineado=0.93, espacio=8)
por_id(24).top = Inches(11.7); por_id(24).height = Inches(5.7)

# --- Franja de 5 pasos ---
encabezado(6.5, 17.6, 13.5, "From literature to experimental decision")
pasos = ["Scientific\nliterature\n(PDFs)", "Automated\nsearch and\nretrieval",
         "Structured\nvariable\nextraction", "Experimental\nDecision Matrix", "Experimental\ndesign criteria"]
px, pw, py, ph = 6.5, 2.40, 18.6, 2.5
for i, s in enumerate(pasos):
    x = px + i * (pw + 0.32)
    panel(x, py, pw, ph, relleno=SUAVE if i < 3 else RGBColor(0xEC, 0xE1, 0xF5))
    caja(x + 0.1, py + 0.25, pw - 0.2, ph - 0.4, s,
         tam=21, bold=(i >= 3), color=MORADO if i >= 3 else TINTA,
         espacio=0, interlineado=0.92, align=PP_ALIGN.CENTER)
    if i < 4:
        caja(x + pw + 0.02, py + 0.8, 0.30, 0.9, "→", tam=26, color=MORADO, bold=True, espacio=0)

encabezado(6.5, 21.5, 13.5, "The Experimental Decision Matrix")
caja(6.5, 22.5, 13.5, 4.4,
     [("The EDM is the central artifact of the project. It is not a descriptive table: it is a decision "
       "instrument. Each row states a decision dimension, what the curated evidence reports, the criterion "
       "adopted for the laboratory, and the numbered sources that support it.", {}),
      ("The design adopted here, aqueous ethanol at 50–70% v/v, balances what the curated corpus reports "
       "against what a teaching laboratory can actually run.", {"bold": True})],
     tam=31, interlineado=0.93, espacio=12)

# --- Fragmento de codigo real ---
encabezado(6.5, 27.1, 13.5, "Why the matrix is auditable")
panel(6.4, 28.1, 13.7, 7.5, relleno=CODIGO, borde=CODIGO)
codigo = (
"// Output schema: the EDM is typed, not prose\n"
"mde: {\n"
"  dimensiones: [{\n"
"    dimension:          string,  // decision variable\n"
"    hallazgo_evidencia: string,  // what the corpus reports\n"
"    criterio_adoptado:  string,  // what the lab will do\n"
"    fuentes: [integer]           // papers backing the row\n"
"  }]\n"
"}\n"
"\n"
"// Defensibility audit: every signal must be anchored\n"
"ancla_tipo: enum [ \"fuente_bibliografica\",\n"
"                   \"ausencia_en_corpus\",\n"
"                   \"resultado_propio\" ]")
caja(6.8, 28.4, 13.0, 6.9, codigo, tam=21, color=RGBColor(0xE8, 0xF0, 0xE3),
     espacio=0, interlineado=1.02, fuente="Consolas")
caja(6.5, 35.9, 13.5, 4.4,
     [("Constraining the model to a schema is what makes the matrix auditable. The fuentes field forces every "
       "adopted criterion back to specific curated papers, and the audit step must anchor each signal either to "
       "a real source or to an explicitly declared absence of evidence.", {}),
      ("The agent cannot answer from intuition: if the corpus does not support a value, the correct output is a "
       "declared gap and a proposed experiment, not a fabricated number.", {"bold": True})],
     tam=27, interlineado=0.93, espacio=10)

encabezado(6.5, 40.4, 13.5, "What this semester closes")
caja(6.5, 41.4, 13.5, 3.2,
     [("Structured extraction of quantitative variables (solvent, method, yield, TPC, IC50/DPPH, conditions) "
       "still lives in a separate assistant. Integrating it into the pipeline is the goal for August to December 2026, "
       "and it is the real ceiling of the system today. Target: EugenIA v2.0.", {})],
     tam=27, interlineado=0.93, espacio=0)

# ---------- COLUMNA DERECHA ----------
escribir(por_id(25), "Results and figures", tam=54, color=MORADO, bold=True, espacio=0)
borrar(26, 27)

CAPS = "capturas"
figs = [("14-mde-tabla.png", "Fig 1. Experimental Decision Matrix"),
        ("07-busqueda-auditable.png", "Fig 2. Auditable search"),
        ("11-bibliografia-aceptadas.png", "Fig 3. Curated bibliography")]
fx, fw, fy = 21.5, 4.25, 4.7
fh = fw * 2532 / 1170
for i, (arch, _) in enumerate(figs):
    ruta = os.path.join(CAPS, arch)
    if os.path.exists(ruta):
        sl.shapes.add_picture(ruta, Inches(fx + i * (fw + 0.35)), Inches(fy), Inches(fw), Inches(fh))

cy = fy + fh + 0.25
caja(21.5, cy, 13.5, 4.4,
     [("Fig 1. EDM generated from a curated corpus. In the first row the system states that no reference on "
       "P. amboinicus reports a systematic TPC comparison across ethanol concentrations, anchors its "
       "recommendation to a study on another species, and proposes an experiment instead of inventing a value.", {}),
      ("Fig 2. Auditable search: the executed query and the records returned by each database. The researcher "
       "can rewrite the query, filter by year and document type, and repeat it.", {}),
      ("Fig 3. Curated bibliography. Each record carries source database, document type, journal, open-access "
       "status and a fit assessment. References enter as pending: the researcher accepts them explicitly.", {})],
     tam=23, color=GRIS, interlineado=0.92, espacio=9)

escribir(por_id(28), "Discussion", tam=54, color=MORADO, bold=True, espacio=0)
por_id(28).top = Inches(19.0)
escribir(por_id(29),
    [("Solvent polarity and extraction method are convergent variables linking phenolic recovery to bioactivity. "
      "The wider bibliometric picture shows accelerating output with an inflection around 2016, and roughly 70% "
      "of production concentrated in the Indo-Malayan region and Brazil, which raises an obvious question for our "
      "own setting: where is the Caribbean?", {}),
     ("The literature also shows a weak link between composition, bioactivity and application. That disconnection "
      "is precisely the gap a decision instrument can help close.", {}),
     ("A limitation worth stating plainly: the EDM is only as good as the abstracts feeding it. Today it is built "
      "from prose, so it can recommend a concentration range but cannot yet assemble a table of medians and study counts.",
      {"color": GRIS, "tam": 28})],
    tam=28, interlineado=0.93, espacio=10)
por_id(29).top = Inches(19.9); por_id(29).height = Inches(6.4)

escribir(por_id(30), "Conclusions", tam=54, color=MORADO, bold=True, espacio=0)
por_id(30).top = Inches(26.6)
escribir(por_id(31),
    [("Curated bibliometrics can be turned into operating criteria, not just a reading list.", {}),
     ("Automated curation still needs verification: 91.3% agreement means about one variable in eleven required human correction.", {}),
     ("The errors were systematic, not random: unit conversion between studies, and ranges reported as point values.", {}),
     ("Declaring an absence of evidence is itself a result. Several EDM rows correctly concluded that no direct "
      "evidence exists for P. amboinicus and proposed an experiment instead.", {})],
    tam=28, interlineado=0.93, espacio=10)
por_id(31).top = Inches(27.5); por_id(31).height = Inches(6.2)

encabezado(21.5, 34.0, 13.5, "Validation of the automated curation")
panel(21.4, 35.0, 13.7, 3.1, relleno=SUAVE)
caja(21.8, 35.15, 4.4, 1.6, "91.3%", tam=76, color=MORADO, bold=True, espacio=0, interlineado=0.9)
caja(26.0, 35.3, 8.8, 2.6, "agreement between EugenIA and manual extraction of experimental variables",
     tam=25, color=GRIS, espacio=0, interlineado=0.92)
caja(21.5, 38.4, 13.5, 4.4,
     [("Errors concentrated in two places: converting units between studies, and reading a reported range as if it "
       "were a point value.", {}),
      ("The practical lesson is that automating the search is not enough. Every value that feeds the matrix has to "
       "be justified and verified before it becomes a design criterion.", {"bold": True})],
     tam=27, interlineado=0.93, espacio=10)

# ---------- PIE ----------
escribir(por_id(33), "References and data availability", tam=42, color=TEAL, bold=True, espacio=0)
escribir(por_id(34),
    "Databases searched: OpenAlex, Europe PMC and Crossref. Open-access enrichment via Unpaywall; PDF validation "
    "with PyMuPDF. The curated 47-article corpus and the full search log are exportable from the tool in BibTeX, "
    "RIS and CSV, and are available from the authors.",
    tam=22, color=GRIS, interlineado=0.92, espacio=0)
por_id(34).height = Inches(1.6)

escribir(por_id(35), "Acknowledgments", tam=42, color=TEAL, bold=True, espacio=0)
escribir(por_id(36),
    "This research is an undergraduate project by Kamila López Díaz at the University of Puerto Rico, Río Piedras "
    "Campus (CIFI 4996), mentored by Prof. Manuel Reyes-Guzmán, Department of Physical Sciences, Faculty of General "
    "Studies. We thank the Puerto Rico Section of the American Chemical Society. EugenIA Lab is named after Eugene "
    "Garfield, founder of modern bibliometrics.",
    tam=22, color=GRIS, interlineado=0.92, espacio=0)
por_id(36).height = Inches(1.6)

prs.save(SALIDA)
print("Guardado:", SALIDA, os.path.getsize(SALIDA), "bytes")
print("Formas totales:", len(sl.shapes))
